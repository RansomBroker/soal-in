import streamlit as st
import math
import io
import re
import requests
from docx import Document
from docx.shared import Inches
from duckduckgo_search import DDGS
from src.config.settings import load_config, set_api_keys
from src.rag.retriever import get_context
from src.generator.qa_generator import generate_questions

st.set_page_config(page_title="Tahap 2: Generate Soal", page_icon="📝", layout="wide")
st.title("📝 Tahap 2: AI Soal & Kisi-kisi Generator")

st.markdown("""
Setelah semua materi terhimpun di Pinecone, rancang parameter ujian yang Anda kehendaki. AI akan membaca ratusan halaman dokumen Anda menggunakan metode MMR dan menuangkannya ke berbagai format tes spesifik (HOTS/LOTS)!
""")

# ---- KUNCI PENGAMANAN HALAMAN ----
if "processed_files" not in st.session_state or not st.session_state.processed_files:
    st.warning("🚫 Stop! Anda belum mengunggah Satupun Materi (PDF) ke Sistem.")
    st.info("Silakan kembali ke halaman **Tahap 1** di menu samping untuk melakukan Setup Knowledge Base terlebih dahulu.")
    st.stop() # Script akan berhenti disini dan tak menampilkan sisa UI bawahnya!

config = load_config()

# ---- 1. SETUP UI ---- #
with st.container(border=True):
    st.subheader("⚙️ Parameter Database (Konteks)")
    top_k_retrieve = st.number_input("Maksimal Penarikan Konteks (Jumlah Paragraf dari Pinecone)", 
                                     min_value=10, max_value=200, value=40,
                                     help="AI butuh bacaan dari PDF untuk membuat soal. Jika Anda mau mencetak 40 soal, angkanya usahakan besar (misal 50 - 80) agar AI punya banyak paragraf acak untuk dikembangkan!")
    topik_soal = st.text_input("Kata Kunci Pencarian Topik", placeholder="Biarkan kosong jika ingin mengambil materi acak dari seluruh PDF.")
    
st.subheader("Distribusi Jumlah Menu Soal")

# Tracker Variabel Global
total_soal = 0
total_hots = 0

cnt_pg, hots_pg_pct = 0, 0
cnt_pgk, hots_pgk_pct = 0, 0
cnt_bs, hots_bs_pct = 0, 0
cnt_pg, hots_pg_pct = 0, 0
cnt_pgk, hots_pgk_pct = 0, 0
cnt_bs, hots_bs_pct = 0, 0
cnt_jd, hots_jd_pct = 0, 0
cnt_us, hots_us_pct = 0, 0 # Uraian Singkat
topik_pg, topik_pgk, topik_bs, topik_jd, topik_us = "", "", "", "", ""
pg_format = "A-E"

# A) PILIHAN GANDA BIASA
use_pg = st.checkbox("Pilihan Ganda (Biasa)")
if use_pg:
    with st.container(border=True):
        c1, c2, c3, c4 = st.columns(4)
        cnt_pg = c1.number_input("Jumlah Soal", 1, 100, 5, key="cnt_pg")
        pg_start = c2.text_input("Opsi Awal", "A", key="pg_start")
        pg_end = c3.text_input("Opsi Akhir", "E", key="pg_end")
        hots_pg_pct = c4.slider("Target % HOTS", 0, 100, 10, step=5, key="hots_pg")
        topik_pg = st.text_input("Topik Khusus (Opsional)", placeholder="Contoh: Struktur Atom", key="t_pg")
        
        pg_format = f"{pg_start.upper()}-{pg_end.upper()}"
        total_soal += cnt_pg
        total_hots += math.ceil(cnt_pg * (hots_pg_pct / 100.0))

# B) PILIHAN GANDA KOMPLEKS
use_pgk = st.checkbox("Pilihan Ganda Kompleks (Lebih dari 1 Jawaban)")
if use_pgk:
    with st.container(border=True):
        c1, c2 = st.columns(2)
        cnt_pgk = c1.number_input("Jumlah Soal", 1, 100, 5, key="cnt_pgk")
        hots_pgk_pct = c2.slider("Target % HOTS", 0, 100, 10, step=5, key="hots_pgk")
        topik_pgk = st.text_input("Topik Khusus (Opsional)", placeholder="Contoh: Hidrokarbon", key="t_pgk")
        
        total_soal += cnt_pgk
        total_hots += math.ceil(cnt_pgk * (hots_pgk_pct / 100.0))

# C) BENAR / SALAH
use_bs = st.checkbox("Benar / Salah")
if use_bs:
    with st.container(border=True):
        c1, c2 = st.columns(2)
        cnt_bs = c1.number_input("Jumlah Soal", 1, 100, 5, key="cnt_bs")
        hots_bs_pct = c2.slider("Target % HOTS", 0, 100, 10, step=5, key="hots_bs")
        topik_bs = st.text_input("Topik Khusus (Opsional)", placeholder="Topik spesifik untuk soal B/S", key="t_bs")
        
        total_soal += cnt_bs
        total_hots += math.ceil(cnt_bs * (hots_bs_pct / 100.0))

# D) MENJODOHKAN
use_jd = st.checkbox("Menjodohkan (Matching)")
if use_jd:
    with st.container(border=True):
        c1, c2 = st.columns(2)
        cnt_jd = c1.number_input("Jumlah Soal", 1, 50, 5, key="cnt_jd")
        hots_jd_pct = c2.slider("Target % HOTS", 0, 100, 10, step=5, key="hots_jd")
        topik_jd = st.text_input("Topik Khusus (Opsional)", placeholder="Topik spesifik untuk Menjodohkan", key="t_jd")
        
        total_soal += cnt_jd
        total_hots += math.ceil(cnt_jd * (hots_jd_pct / 100.0))

# E) URAIAN SINGKAT (Baru)
use_us = st.checkbox("Uraian Singkat / Hitungan Esai")
if use_us:
    with st.container(border=True):
        c1, c2 = st.columns(2)
        cnt_us = c1.number_input("Jumlah Soal", 1, 50, 5, key="cnt_us")
        hots_us_pct = c2.slider("Target % HOTS", 0, 100, 10, step=5, key="hots_us")
        topik_us = st.text_input("Topik Khusus (Opsional)", placeholder="Fokus materi hitungan/esai", key="t_us")
        
        total_soal += cnt_us
        total_hots += math.ceil(cnt_us * (hots_us_pct / 100.0))


total_lots = total_soal - total_hots

# Highlight Rekapitulasi
st.info(f"**REKAPITULASI APLIKASI:** \n\n"
        f"Total Rancangan Soal: **{total_soal} Butir**\n"
        f"🎯 Terakumulasi total **{total_hots} Soal HOTS** dan **{total_lots} Soal LOTS** secara keseluruhan.")

# Variabel State untuk Previewing
if "generated_results" not in st.session_state:
    st.session_state["generated_results"] = {}

# ---- 2. GENERATION ACTION ---- #
if st.button("🚀 Mulai Generate Bank Soal", type="primary", use_container_width=True):
    if total_soal == 0:
        st.warning("Peringatan: Anda belum menceklis satupun bentuk soal. Silakan pilih jenis soal yang dimau.")
    elif not config["google_api_key"] or not config["pinecone_api_key"]:
        st.error("🔑 API Key belum tersedia. Pastikan konfigurasi aman.")
    else:
        set_api_keys(config["google_api_key"], config["pinecone_api_key"])
        st.session_state["generated_results"] = {} # Kosongkan preview lama
        
        with st.status("Mesin AI Sedang Bekerja Membongkar File Database...", expanded=True) as status_box:
            try:
                st.write("🔍 Merangkum literatur referensi dengan algoritma pencarian MMR-Pinecone...")
                
                # Topik dasar dari parameter utama (Bila yang spesifik kosong)
                query_topik = topik_soal if topik_soal.strip() else "Materi penting"
                
                # Fungsi Helper untuk mengeksekusi AI per Tipe
                def trigger_ai(tipe_soal, jumlah, param_hots_pct, topik_tipe="", op_length="A-E"):
                    if jumlah <= 0: return
                    
                    # Tentukan Topik Final: Cek apakah user input topik khusus, jika tidak pakai topik global
                    final_topik = topik_tipe if topik_tipe.strip() else query_topik
                    
                    st.write(f"💭 Mencetak {jumlah} Soal tipe {tipe_soal} Topik: '{final_topik}'...")
                    
                    # Cari referensi ke Pinecone SECARA SPESIFIK untuk topik ini!
                    konteks_kasar = get_context(final_topik, config["pinecone_index_name"], top_k=top_k_retrieve)
                    
                    jml_hots = math.ceil(jumlah * (param_hots_pct / 100.0))
                    jml_lots = jumlah - jml_hots
                    
                    hasil = generate_questions(
                        topic=final_topik,
                        context_text=konteks_kasar,
                        item_count=jumlah,
                        hots_count=jml_hots,
                        lots_count=jml_lots,
                        question_type=tipe_soal,
                        pg_options=op_length
                    )
                    
                    st.session_state["generated_results"][tipe_soal] = hasil
                
                # Jalankan Siklus Sesuai Ceklis User
                if use_pg:
                    trigger_ai("Pilihan Ganda", cnt_pg, hots_pg_pct, topik_pg, pg_format)
                if use_pgk:
                    trigger_ai("Pilihan Ganda Kompleks", cnt_pgk, hots_pgk_pct, topik_pgk)
                if use_bs:
                    trigger_ai("Benar/Salah", cnt_bs, hots_bs_pct, topik_bs)
                if use_jd:
                    trigger_ai("Menjodohkan", cnt_jd, hots_jd_pct, topik_jd)
                if use_us:
                    trigger_ai("Uraian Singkat", cnt_us, hots_us_pct, topik_us)
                
                status_box.update(label="Seluruh Soal Telah Berhasil Dicetak ke Layar Anda!", state="complete", expanded=False)
                
            except Exception as e:
                status_box.update(label="Terjadi Gangguan pada Koneksi/Batas Quota!", state="error")
                st.error(str(e))

# ---- 3. WEB PREVIEW SECTION ---- #
if len(st.session_state["generated_results"]) > 0:
    st.divider()
    st.header("📖 Preview & Export Hasil")
    
    import re
    
    # Fungsi Helper Export ke Word (DOCX) + Parser LaTeX & Table
    # Fungsi Helper Export ke Word (DOCX) + Parser LaTeX & Table
    def create_docx(text_content):
        doc = Document()
        
        in_table = False
        table_obj = None
        
        def apply_tokens_to_paragraph(p, text):
            # Memecah kalimat berdasarkan pola Bold (**), Latex Math ($), dan Gambar (![])
            tokens = re.split(r'(\*\*.*?\*\*|\$.*?\$|!\[.*?\]\(.*?\))', text)
            for token in tokens:
                if not token:
                    continue
                if token.startswith('**') and token.endswith('**'):
                    p.add_run(token[2:-2]).bold = True
                elif token.startswith('![') and '](' in token and token.endswith(')'):
                    url_start = token.find('](') + 2
                    url = token[url_start:-1]
                    try:
                        resp = requests.get(url, timeout=4)
                        if resp.status_code == 200:
                            img_stream = io.BytesIO(resp.content)
                            p.add_run().add_picture(img_stream, width=Inches(3.5))
                        else:
                            p.add_run(f"\n[Visualisasi Hilang: {url}]\n").italic = True
                    except Exception:
                        p.add_run(f"\n[Gagal mengunduh Visualisasi: {url}]\n").italic = True
                elif token.startswith('$') and token.endswith('$'):
                    math_text = token[1:-1].replace('\\rightarrow', ' → ')
                    sub_tokens = re.split(r'(_[a-zA-Z0-9]|\^[a-zA-Z0-9]|_{[^}]+}|\^{[^{]+})', math_text)
                    for sub in sub_tokens:
                        if sub.startswith('_'):
                            val = sub[1:]
                            if val.startswith('{') and val.endswith('}'): val = val[1:-1]
                            p.add_run(val).font.subscript = True
                        elif sub.startswith('^'):
                            val = sub[1:]
                            if val.startswith('{') and val.endswith('}'): val = val[1:-1]
                            p.add_run(val).font.superscript = True
                        else:
                            p.add_run(sub)
                else:
                    p.add_run(token)
        
        for line in text_content.split('\n'):
            line = line.strip()
            if not line:
                if in_table:
                    in_table = False
                    table_obj = None
                doc.add_paragraph()
                continue
                
            if line == "---PAGE_BREAK---":
                if in_table:
                    in_table = False
                    table_obj = None
                doc.add_page_break()
                continue
            
            # Deteksi Markdown Table
            if line.startswith('|') and line.endswith('|'):
                # Abaikan baris pemisah header |---|---|
                if re.match(r'^\|[\s\-\|]+\|$', line):
                    continue
                
                cells = [c.strip() for c in line.split('|')[1:-1]]
                
                if not in_table:
                    in_table = True
                    table_obj = doc.add_table(rows=1, cols=len(cells))
                    table_obj.style = 'Table Grid'
                    # Isi Header
                    hdr_cells = table_obj.rows[0].cells
                    for i, text in enumerate(cells):
                        if i < len(hdr_cells):
                            # Make header bold
                            p = hdr_cells[i].paragraphs[0]
                            clean_text = re.sub(r'(?i)<br\s*/?>', '\n', text)
                            apply_tokens_to_paragraph(p, f"**{clean_text.replace('**', '')}**")
                else:
                    # Tambah baris data
                    row_cells = table_obj.add_row().cells
                    for i, text in enumerate(cells):
                        if i < len(row_cells):
                            # Render <br> menjadi native line break di MS Word
                            clean_text = re.sub(r'(?i)<br\s*/?>', '\n', text)
                            p = row_cells[i].paragraphs[0]
                            apply_tokens_to_paragraph(p, clean_text)
                continue
            else:
                if in_table:
                    in_table = False
                    table_obj = None
                    
            if line.startswith('### ') or line.startswith('## '):
                doc.add_heading(line.replace('### ', '').replace('## ', ''), level=3)
            elif line.startswith('# '):
                doc.add_heading(line.replace('# ', ''), level=1)
            else:
                p = doc.add_paragraph()
                
                # Pembersih Opsi A-Z Super Kuat (Menghapus minus, asterisk, atau spasi berlebih)
                if re.match(r'^[\-\*]\s*([A-Z]\.)', line):
                    # Match "- A.", "* B.", "-  C."
                    line = re.sub(r'^[\-\*]\s*', '', line)
                    p.paragraph_format.left_indent = Inches(0.2)
                elif re.match(r'^([A-Z]\.)', line):
                    # Terkadang AI langsung menulis "A." tanpa bullet awalan
                    p.paragraph_format.left_indent = Inches(0.2)
                elif re.match(r'^[\-\*]\s*\[[\sXx]\]\s*([A-Z]\.)', line):
                    # Jika AI bandel menulis "- [ ] A."
                    line = re.sub(r'^[\-\*]\s*\[[\sXx]\]\s*', '', line)
                    p.paragraph_format.left_indent = Inches(0.2)
                elif line.startswith('- ') or line.startswith('* '):
                    # List biasa (misal poin-poin pertanyaan)
                    p.style = 'List Bullet'
                    line = line[2:]
                
                apply_tokens_to_paragraph(p, line)
        
        bio = io.BytesIO()
        doc.save(bio)
        return bio.getvalue()
    
    # Bundling semua teks untuk fitur Download Master dengan Struktur Ujian Resmi
    all_soal = []
    all_kunci = []
    all_kisi = []
    
    for k, v in st.session_state["generated_results"].items():
        # Memisahkan bagian (### SOAL), (### KUNCI), dan (### KISI) via Multiline Regex penuh
        parts = re.split(r'^(### .*?)$', v, flags=re.MULTILINE)
        
        current_header = None
        for p in parts:
            p_strip = p.strip()
            if not p_strip: continue
            
            if 'SOAL' in p_strip and 'KISI' not in p_strip:
                current_header = 'SOAL'
                all_soal.append(f"\n**Bagian: {k.upper()}**\n")
            elif 'KUNCI' in p_strip:
                current_header = 'KUNCI'
                all_kunci.append(f"\n**Bagian: {k.upper()}**\n")
            elif 'KISI' in p_strip:
                current_header = 'KISI'
            else:
                if current_header == 'SOAL':
                    all_soal.append(p)
                elif current_header == 'KUNCI':
                    all_kunci.append(p)
                elif current_header == 'KISI':
                    all_kisi.append(p)

    master_text = "# LEMBAR SOAL UJIAN\n" + "".join(all_soal) + "\n\n---PAGE_BREAK---\n\n" + "# KUNCI JAWABAN\n" + "".join(all_kunci) + "\n\n---PAGE_BREAK---\n\n" + "# KISI-KISI SOAL\n" + "".join(all_kisi)
    
    st.download_button(
        label="📥 Download Semua Soal & Kisi-Kisi (.DOCX)",
        data=create_docx(master_text),
        file_name="Bank_Soal_Generated.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )
    
    # Gunakan Tabs agar preview sangat rapi per Jenis Soal
    tabs = st.tabs(list(st.session_state["generated_results"].keys()))
    
    for tab, (tipe_soal, teks_hasil) in zip(tabs, st.session_state["generated_results"].items()):
        with tab:
            st.download_button(
                label=f"📥 Download Spesifik: {tipe_soal} (.DOCX)",
                data=create_docx(teks_hasil),
                file_name=f"Soal_{tipe_soal.replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"dl_btn_{tipe_soal}"
            )
            st.markdown(teks_hasil, unsafe_allow_html=True)
            st.caption(f"*Soal, Kisi-kisi, dan Kemenjodohan di atas diproduksi secara eksklusif menggunakan sistem AI berdasarkan dokumen milik Anda pribadi.*")
